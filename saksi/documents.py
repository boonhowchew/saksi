# SPDX-License-Identifier: AGPL-3.0-only
# Copyright (c) 2026 Boon-How Chew

"""Plain-text extraction from protocol documents (txt, md, pdf, docx)."""

from __future__ import annotations

import io
import re
import unicodedata
from collections import Counter
from difflib import SequenceMatcher
from typing import Any


class UnsupportedDocument(ValueError):
    pass


def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="replace")
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    raise UnsupportedDocument(f"cannot read {filename}; use .txt, .md, .pdf or .docx")


def normalise(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)        # ligatures (fi), full-width and compatibility forms
    for bad in ("\u00ad", "\u200b", "\u200c", "\u200d", "\ufeff"):   # soft hyphen, zero-width characters
        text = text.replace(bad, "")
    text = text.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    text = re.sub(r"[\u2010-\u2015\u2212]", "-", text)    # dashes and the minus sign
    text = re.sub(r"[\u00a0\u2000-\u200a\u202f\u205f]", " ", text)   # non-breaking and thin spaces
    text = re.sub(r"-\s*\n\s*", "", text)          # rejoin words hyphenated across lines (PDFs)
    return re.sub(r"\s+", " ", text).strip().lower()


_TOKEN = re.compile(r"[a-z0-9]+")
CLOSE_MATCH = 0.9          # share of the quotation's words that must appear, in order, in the protocol


def locate(quote: str, source: str) -> dict[str, Any]:
    """Find a quotation in a source document.

    "exact": found after normalising spacing, punctuation variants and line breaks.
    "close": at least 90% of its words appear in order in one passage of similar length (PDF text
             often carries ligatures, footers or reference numbers mid-sentence), AND every number in
             the quotation appears in that passage. The passage's own wording is returned.
    "not_found": otherwise. Too-short quotations are never found."""
    q = normalise(quote or "").strip(" .\"'")
    if len(q) < 12:
        return {"status": "not_found", "why": "too short to check" if q else "not given"}
    src = normalise(source or "")
    if q in src:
        return {"status": "exact", "protocol_text": q}
    qt = _TOKEN.findall(q)
    spans = list(_TOKEN.finditer(src))
    st = [m.group(0) for m in spans]
    if len(qt) < 4 or not st:
        return {"status": "not_found", "why": "not in the protocol"}
    freq = Counter(st)
    anchors = sorted({t for t in qt if freq.get(t)}, key=lambda t: (freq[t], -len(t)))[:3]
    L, slack = len(qt), max(4, len(qt) // 2)
    best_n, best_span = 0, None
    for a in anchors:
        qi = qt.index(a)
        for pos in (i for i, t in enumerate(st) if t == a):
            start = max(0, pos - qi - slack // 2)
            window = st[start:start + L + slack]
            blocks = [b for b in SequenceMatcher(None, qt, window, autojunk=False).get_matching_blocks() if b.size]
            n = sum(b.size for b in blocks)
            if n > best_n:
                best_n, best_span = n, (start + blocks[0].b, start + blocks[-1].b + blocks[-1].size)
    if best_span and best_n / L >= CLOSE_MATCH:
        i, j = best_span
        passage = src[spans[i].start():spans[j - 1].end()]
        numbers = [t for t in qt if t.isdigit()]
        if all(n in set(_TOKEN.findall(passage)) for n in numbers):
            return {"status": "close", "protocol_text": passage, "match": round(best_n / L, 3)}
        return {"status": "not_found", "why": "a similar passage exists, but its numbers differ", "protocol_text": passage}
    return {"status": "not_found", "why": "not in the protocol"}


def quote_found(quote: str, source: str) -> bool:
    return locate(quote, source)["status"] != "not_found"
