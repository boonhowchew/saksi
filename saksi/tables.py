# SPDX-License-Identifier: AGPL-3.0-only
# Copyright (c) 2026 Boon-How Chew

"""Publication-ready tables, built by code from the recorded results. No model touches them.

Table 1 (characteristics, stratified) comes from the standard step; every other analysis with
estimates gets a results table, crude and adjusted side by side when both are reported. Tables are
numbered in reporting order (the standard sequence), the same numbers the Writing Agent is given.
Formatting follows common journal practice: three horizontal rules, no vertical lines, estimates
with 95% CI as "x (a to b)", P values to three decimals or "<0.001", footnotes for tests,
abbreviations, the sample and the label (confirmatory or exploratory).

The page shows the same table objects; `to_docx` gives a Word file the authors can paste from.
"""

from __future__ import annotations

import io
import json
import re
from typing import Any

from .blobs import BlobStore
from .standard import position

ABBREVIATIONS = [("odds ratio", "OR"), ("hazard ratio", "HR"), ("risk ratio", "RR"), ("relative risk", "RR"),
                 ("incidence rate ratio", "IRR"), ("rate ratio", "IRR"), ("mean difference", "MD"),
                 ("area under the curve", "AUC"), ("area under the roc curve", "AUC")]
LETTERS = "abcdefghij"


def _p(p: Any) -> str:
    if not isinstance(p, (int, float)) or isinstance(p, bool):
        return ""
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def _x(v: Any, d: int | None = None) -> str:
    if not isinstance(v, (int, float)) or isinstance(v, bool):
        return "" if v is None else str(v)
    if d is None:
        d = 3 if v != 0 and abs(v) < 0.1 else 2
    return f"{v:.{d}f}"


def _measure(m: str) -> tuple[str, list[tuple[str, str]]]:
    """'adjusted odds ratio' -> ('Adjusted OR', [('OR', 'odds ratio')])."""
    low, used = (m or "").strip().lower(), []
    out = m.strip() if m else "Estimate"
    for full, short in ABBREVIATIONS:
        if full in low:
            out = re.sub(full, short, out, flags=re.I)
            used.append((short, full))
            break
    if low in ("auc", "c-statistic"):
        used.append(("AUC", "area under the receiver operating characteristic curve"))
    return out[:1].upper() + out[1:], used


def _labels(dictionary: list[dict[str, Any]] | None) -> dict[str, dict[str, Any]]:
    return {d["column"]: d for d in dictionary or [] if isinstance(d, dict) and d.get("column")}


def _readable(col: str) -> str:
    """A column name with no dictionary label, made readable: 'previous_gdm' -> 'Previous GDM'."""
    words = re.split(r"[_.\s]+", str(col).strip())
    out = [w.upper() if w.lower() in ACRONYMS else w for w in words if w]
    text = " ".join(out) or str(col)
    return text[:1].upper() + text[1:]


ACRONYMS = {"bmi", "gdm", "hba1c", "hb", "sbp", "dbp", "bp", "ldl", "hdl", "egfr", "ogtt", "fbs", "fpg", "id", "hiv",
            "icu", "auc", "dm", "t2dm", "ckd", "ses", "ga", "crp", "alt", "ast", "tg", "tc"}


def _name(col: str, labels: dict[str, dict[str, Any]], continuous: bool = False) -> str:
    d = labels.get(col) or {}
    name = str(d.get("label") or _readable(col))
    units = str(d.get("units_or_coding") or "").strip()
    if continuous and _is_unit(units):
        name += f", {units}"
    return name


def _is_unit(units: str) -> bool:
    """A unit worth printing ('years', 'kg/m2', 'g/dL'), not a coding, range, count or a note to the authors."""
    return bool(units) and len(units) <= 16 and not re.search(
        r"\d\s*/\s*\d|=|\d\s*(?:-|–|to)\s*\d|unknown|confirm|unclear|\?|count|code|categor|scale|score|none|n/a|—", units, re.I)


def _label_note(a: dict[str, Any]) -> str:
    if a.get("label") == "confirmatory":
        return ("Primary (confirmatory) analysis; named primary by the authors, as stated in Methods."
                if a.get("primary_basis") == "authors" else "Primary (confirmatory) analysis, as pre-specified.")
    return f"Exploratory analysis ({a.get('basis') or 'pre-specified'})."


NA_CODES = {"#NULL!", ".", "NA", "N/A", "nan", "NaN", "NULL"}
_YES, _NO = {"1", "yes", "y", "true"}, {"0", "no", "n", "false"}


def _clean_label(name: str) -> str:
    """'Gestational diabetes yes/no' -> 'Gestational diabetes'."""
    return re.sub(r"[\s,;:(]*\b(?:yes\s*/\s*no|y\s*/\s*n|1\s*/\s*0|0\s*/\s*1)\b\)?\s*$", "", name, flags=re.I).strip() or name


def _codes(col: str, labels: dict[str, dict[str, Any]]) -> dict[str, str]:
    """Value labels from the data dictionary's coding, for example '0=No, 1=Yes' or '1 = Malay; 2 = Chinese'."""
    units = str((labels.get(col) or {}).get("units_or_coding") or "")
    return {k.strip(): v.strip().rstrip(".") for k, v in re.findall(r"(-?\w+)\s*[=:]\s*([^,;=]+?)(?=\s*(?:[,;]|\s-?\w+\s*[=:]|$))", units)}


def _is_binary(levels: list[str]) -> bool:
    low = {str(x).strip().lower() for x in levels}
    return bool(low) and low <= (_YES | _NO) and bool(low & _YES)


def _table1(aid: str, a: dict[str, Any], t: dict[str, Any], labels: dict[str, dict[str, Any]]) -> dict[str, Any]:
    by = t.get("stratify_by")
    groups = [g for g in t.get("groups") or [] if str(g["level"]) not in NA_CODES]
    by_name = _clean_label(_name(by, labels)) if by else ""
    codes_by = _codes(by, labels) if by else {}

    def group_head(level: str) -> str:
        if level in codes_by and codes_by[level].lower() not in _YES | _NO:
            return codes_by[level]
        if _is_binary([g["level"] for g in groups]):
            yes = str(level).strip().lower() in _YES
            return by_name if yes else f"No {by_name[:1].lower()}{by_name[1:]}"
        return f"{by_name}: {codes_by.get(level, level)}"
    cols = ["Characteristic", f"Overall\n(n = {sum(g['n'] for g in groups) if groups else t.get('n_analysed')})"] + \
        [f"{group_head(str(g['level']))}\n(n = {g['n']})" for g in groups] + (["P value"] if groups else [])
    tests: dict[str, str] = {}

    def mark(r: dict[str, Any]) -> str:
        key = r.get("test") or ""
        if not key:
            return ""
        if r.get("test_note"):
            key += " (more than 20% of expected counts below 5; read with care)"
        return tests.setdefault(key, LETTERS[len(tests) % len(LETTERS)])

    def cont(x: dict[str, Any] | None, summary: str) -> str:
        if not x or not x.get("n"):
            return ""
        if summary == "mean (SD)":
            return f"{_x(x.get('mean'), 1)} ({_x(x.get('sd'), 1)})"
        return f"{_x(x.get('median'), 1)} ({_x(x.get('q1'), 1)} to {_x(x.get('q3'), 1)})"

    def npct(x: dict[str, Any] | None) -> str:
        return f"{x['n']} ({_x(x.get('pct'), 1)})" if x and x.get("pct") is not None else (str(x["n"]) if x else "")

    rows, missing = [], []
    for r in t.get("rows") or []:
        if r.get("error"):
            continue
        v = r["variable"]
        if r.get("missing"):
            missing.append(f"{_clean_label(_name(v, labels))} {r['missing']}")
        p = [_p(r.get("p_value"))] if groups else []
        m = {len(cols) - 1: mark(r)} if groups and r.get("p_value") is not None else {}
        if r.get("type") == "continuous":
            s = r.get("summary", "")
            rows.append({"cells": [f"{_name(v, labels, True)}, {s}", cont(r.get("overall"), s)]
                         + [cont((r.get("by_group") or {}).get(g["level"]), s) for g in groups] + p, "marks": m})
            continue
        cats = [c for c in r.get("categories") or [] if str(c["level"]) not in NA_CODES]
        if _is_binary([c["level"] for c in cats]):
            # A yes/no characteristic is one row: the number (%) with it, as journals report it.
            yes = next(c for c in cats if str(c["level"]).strip().lower() in _YES)
            rows.append({"cells": [f"{_clean_label(_name(v, labels))}, n (%)", npct(yes.get("overall"))]
                         + [npct((yes.get("by_group") or {}).get(g["level"])) for g in groups] + p, "marks": m})
            continue
        codes = _codes(v, labels)
        rows.append({"cells": [f"{_clean_label(_name(v, labels))}, n (%)", ""] + [""] * len(groups) + p, "marks": m})
        for c in cats:
            rows.append({"cells": [codes.get(str(c["level"]), str(c["level"])), npct(c.get("overall"))]
                         + [npct((c.get("by_group") or {}).get(g["level"])) for g in groups]
                         + ([""] if groups else []), "indent": True})
    notes = ["Values are n (%), mean (SD) or median (IQR, as lower to upper quartile), chosen for each variable by a "
             "fixed rule from its distribution; for a yes/no characteristic, the number (%) with it."]
    notes += [f"{letter} {test}." for test, letter in tests.items()]
    if missing:
        notes.append("Missing values: " + "; ".join(missing) + ".")
    if t.get("n_excluded"):
        notes.append(f"{t['n_excluded']} records without {by_name.lower()} are not shown.")
    return {"analysis_id": aid, "kind": "table1",
            "title": "Characteristics of the participants" + (f", by {by_name.lower()}" if by else ""),
            "columns": cols, "rows": rows, "footnotes": notes,
            "stale_missing_codes": any(str(g["level"]) in NA_CODES for g in t.get("groups") or [])}


MAX_ROWS, MAX_COLS = 80, 12


def _code_table(raw: Any) -> dict[str, Any] | None:
    """A table laid out by the analysis code (when an author asks for a layout). Every cell is text the code
    formatted from values it computed; Saksi only checks the shape."""
    if not isinstance(raw, dict):
        return None
    cols = raw.get("columns")
    if not isinstance(cols, list) or not 2 <= len(cols) <= MAX_COLS:
        return None
    rows = []
    for r in (raw.get("rows") or [])[:MAX_ROWS]:
        cells = r.get("cells") if isinstance(r, dict) else r
        if not isinstance(cells, list) or len(cells) != len(cols):
            return None
        rows.append({"cells": ["" if c is None else str(c) for c in cells],
                     "indent": bool(r.get("indent")) if isinstance(r, dict) else False})
    if not rows:
        return None
    return {"title": str(raw.get("title") or ""), "columns": [str(c) for c in cols], "rows": rows,
            "footnotes": [str(n) for n in (raw.get("footnotes") or [])[:10]]}


def _common_notes(a: dict[str, Any], out: dict[str, Any]) -> list[str]:
    notes = []
    n, ev, ex = out.get("n_analysed"), out.get("n_events"), out.get("n_excluded")
    if n is not None:
        notes.append(f"n = {n} analysed" + (f", {ev} with the outcome" if ev is not None else "")
                     + (f"; {ex} excluded" + (f" ({out['exclusion_reason']})" if out.get("exclusion_reason") else "")
                        if ex else "") + ".")
    model = (out.get("model") or {}).get("formula_or_description")
    if model:
        notes.append(f"Model: {model}.")
    return notes


# Results that describe groups side by side ("Age mean [GDM]", "Smoking=1 proportion [No GDM]", "Sample size GDM")
# are laid out the way journals print a characteristics table: groups as columns with their n, one row per
# characteristic (a yes/no one as a single row, categories indented under their variable), one P value per row.
_GROUPED = re.compile(r"^(?P<var>.+?)(?:[\s_]+(?P<stat>mean|median|proportion|percent(?:age)?|prevalence|rate|count|n|sd))?"
                      r"\s*[\[(](?P<group>[^\])]+)[\])]\s*$", re.I)
_SIZE = re.compile(r"^(?:sample\s+size|group\s+size|number(?:\s+of\s+participants)?|n|count|total)\b[\s:,=_-]*(?:in\s+)?(?P<group>.*)$", re.I)
_INSTRUCTION = re.compile(r"^\s*(?:please\s+)?(?:analy[sz]e|generate|create|make|produce|compute|calculate|run|show|do|add|"
                          r"perform|build|stratify|compare|estimate|fit|redo|repeat|plot|draw|test|check|tabulate|describe|"
                          r"summari[sz]e|give|report|present|prepare|construct|i\s+want|can\s+you|could\s+you)\b", re.I)


def _dp(v: Any) -> int:
    return 1 if isinstance(v, (int, float)) and abs(v) >= 10 else 2


def is_instruction(title: str) -> bool:
    """A request typed by an author ("Analyse and generate ...") is not a table title."""
    return bool(_INSTRUCTION.match(title or ""))


def _title_from_request(title: str) -> str:
    first = re.split(r"(?<=[.!?])\s+", (title or "").strip())[0].rstrip(".!? ")
    first = re.sub(r"^(?:(?:please|and|then|also)\s+|(?:analy[sz]e|generate|create|make|produce|compute|calculate|run|show|"
                   r"do|add|perform|build|compare|estimate|fit|plot|draw|tabulate|describe|summari[sz]e|give|report|present|"
                   r"prepare|construct)\s+(?:and\s+)?|(?:a|an|the)\s+)+", "", first, flags=re.I).strip()
    return (first[:1].upper() + first[1:]) if first else "Results"


def _grouped(aid: str, a: dict[str, Any], out: dict[str, Any], est: list[dict[str, Any]],
             labels: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
    groups: list[str] = []
    sizes: dict[str, Any] = {}
    data: dict[str, dict[str, dict[str, Any]]] = {}     # column -> level ("" for continuous) -> group -> entry
    order: list[str] = []
    for e in est:
        name = str(e["name"]).strip()
        m = _GROUPED.match(name)
        var = (m["var"] if m else name).strip()
        sm = _SIZE.match(var) if m else _SIZE.match(name)
        if sm and isinstance(e.get("estimate"), (int, float)) and (m or sm["group"].strip()):
            g = (m["group"] if m else sm["group"]).strip()
            sizes[g] = e["estimate"]
            if g not in groups:
                groups.append(g)
            continue
        if not m:
            return None
        g = m["group"].strip()
        stat = (m["stat"] or str(e.get("measure") or "")).lower()
        col, _, level = var.partition("=")
        col, level = col.strip(), level.strip()
        kind = ("pct" if re.search(r"proportion|percent|prevalence|%", stat) else "median" if "median" in stat
                else "n" if re.fullmatch(r"count|n|number", stat.strip()) else "mean")
        if g not in groups:
            groups.append(g)
        if col not in data:
            data[col], _ = {}, order.append(col)
        data[col].setdefault(level, {})[g] = e | {"_kind": kind}
    if len(groups) < 2 or len(data) < 2 or len(groups) > MAX_COLS - 2:
        return None

    def cell(e: dict[str, Any] | None) -> str:
        if not e or not isinstance(e.get("estimate"), (int, float)):
            return ""
        v, lo, hi, k = e["estimate"], e.get("ci_lower"), e.get("ci_upper"), e["_kind"]
        scale = 100 if k == "pct" and v <= 1 and (hi is None or hi <= 1) else 1
        if k == "n":
            return f"{v:.0f}"
        d = 1 if k == "pct" else _dp(v)
        ci = f" ({lo * scale:.{d}f} to {hi * scale:.{d}f})" if isinstance(lo, (int, float)) and isinstance(hi, (int, float)) else ""
        if e.get("sd") is not None and not ci:
            ci = f" ({_x(e['sd'], d)})"
        return f"{v * scale:.{d}f}{ci}"

    def stat_label(entries: list[dict[str, Any]]) -> str:
        e = entries[0]
        has_ci = any(x.get("ci_lower") is not None for x in entries)
        base = {"pct": "%", "median": "median", "n": "n", "mean": "mean"}[e["_kind"]]
        return f"{base} (95% CI)" if has_ci and base != "n" else (f"{base} (SD)" if any(x.get("sd") is not None for x in entries) else base)

    def pval(levels: dict[str, dict[str, Any]]) -> str:
        for per in levels.values():
            for x in per.values():
                if x.get("p_value") is not None:
                    return _p(x["p_value"])
        return ""

    rows, kinds = [], set()
    for col in order:
        levels = data[col]
        entries = [x for per in levels.values() for x in per.values()]
        kinds.add(stat_label(entries))
        label = _clean_label(_name(col, labels, continuous=entries[0]["_kind"] in ("mean", "median")))
        codes = _codes(col, labels)
        p = pval(levels)
        keys = list(levels)
        if keys == [""] or (len(keys) == 1 and keys[0].lower() in _YES):
            per = levels[keys[0]]
            rows.append({"cells": [f"{label}, {stat_label(entries)}"] + [cell(per.get(g)) for g in groups] + [p]})
            continue
        rows.append({"cells": [f"{label}, {stat_label(entries)}"] + [""] * len(groups) + [p]})
        for lv in keys:
            name = codes.get(lv) or (lv.capitalize() if lv.isupper() and len(lv) > 4 else lv)
            rows.append({"cells": [name] + [cell(levels[lv].get(g)) for g in groups] + [""], "indent": True})
    has_p = any(r["cells"][-1] for r in rows)
    cols = ["Characteristic"] + [f"{g}\n(n = {sizes[g]:.0f})" if isinstance(sizes.get(g), (int, float)) else g for g in groups]
    if has_p:
        cols.append("P value")
    else:
        for r in rows:
            r["cells"] = r["cells"][:-1]
    said = [("percentage" + k[1:]) if k.startswith("%") else k for k in sorted(kinds)]
    notes = ["Values are " + (", ".join(said[:-1]) + " or " + said[-1] if len(said) > 1 else said[0]) + ", as computed by the analysis code."]
    model = (out.get("model") or {}).get("formula_or_description")
    if model and has_p:
        notes.append(f"P values: {model}.")
    notes.append("CI, confidence interval.")
    spec = a.get("spec") or {}
    oc = spec.get("outcome") if isinstance(spec.get("outcome"), dict) else {}
    by = spec.get("stratify_by") or oc.get("column")
    by_name = _clean_label(_name(by, labels)) if by else ""
    title = f"Characteristics of the participants, by {by_name[:1].lower()}{by_name[1:]}" if by_name else \
        "Characteristics of the participants: " + ", ".join(groups[:-1]) + " and " + groups[-1]
    return {"analysis_id": aid, "kind": "results", "title": title, "columns": cols, "rows": rows,
            "footnotes": notes, "grouped": {"groups": groups, "n": [sizes.get(g) for g in groups], "columns": order}}


def _results(aid: str, a: dict[str, Any], out: dict[str, Any], labels: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
    laid = _code_table(out.get("table"))
    if laid:
        return {"analysis_id": aid, "kind": "results", "title": laid["title"] or str(a["spec"].get("title") or aid),
                "columns": laid["columns"], "rows": laid["rows"],
                "footnotes": laid["footnotes"] + _common_notes(a, out), "laid_out_by": "analysis code"}
    est = [e for e in out.get("estimates") or [] if isinstance(e, dict) and e.get("name") is not None]
    if not est:
        return None
    wide = _grouped(aid, a, out, est, labels)
    if wide:
        return wide
    measures = list(dict.fromkeys(str(e.get("measure") or "Estimate") for e in est))
    terms = list(dict.fromkeys(str(e["name"]) for e in est))
    by = {(str(e["name"]), str(e.get("measure") or "Estimate")): e for e in est}
    cols, abbr = ["Variable"], {"CI": "confidence interval"}
    for m in measures:
        head, used = _measure(m)
        abbr.update(dict(used))
        cols += [f"{head} (95% CI)", "P value"]

    def cell(e: dict[str, Any] | None) -> list[str]:
        if not e:
            return ["", ""]
        lo, hi = e.get("ci_lower"), e.get("ci_upper")
        ci = f" ({_x(lo)} to {_x(hi)})" if lo is not None and hi is not None else ""
        v = e.get("estimate")
        if isinstance(v, float) and v.is_integer() and not ci and re.fullmatch(r"(?i)\s*(count|n|number|sample size|events?)\s*", str(e.get("measure") or "")):
            v = int(v)
        return [f"{_x(v)}{ci}", _p(e.get("p_value"))]

    rows = [{"cells": [_clean_label(_name(t, labels))] + [x for m in measures for x in cell(by.get((t, m)))]}
            for t in terms]
    # Drop P-value columns that are empty throughout (for example an AUC).
    keep = [i for i in range(len(cols)) if cols[i] != "P value" or any(r["cells"][i] for r in rows)]
    cols = [cols[i] for i in keep]
    for r in rows:
        r["cells"] = [r["cells"][i] for i in keep]
    notes = []
    n, ev, ex = out.get("n_analysed"), out.get("n_events"), out.get("n_excluded")
    if n is not None:
        notes.append(f"n = {n} analysed" + (f", {ev} with the outcome" if ev is not None else "")
                     + (f"; {ex} excluded" + (f" ({out['exclusion_reason']})" if out.get("exclusion_reason") else "")
                        if ex else "") + ".")
    model = (out.get("model") or {}).get("formula_or_description")
    if model:
        notes.append(f"Model: {model}.")
    sc = a.get("sample_check") or {}
    if sc.get("shortfall"):
        parts = [f"{r['achieved']} of {r['target']} planned {'participants' if k == 'n' else 'outcome events'}"
                 for k, r in sc["targets"].items() if r.get("short")]
        notes.append("Below the protocol's planned sample: " + "; ".join(parts) + ".")
    d = out.get("diagnostics") if isinstance(out.get("diagnostics"), dict) else None
    if d:
        f, asm = d.get("fit") or {}, d.get("assumptions") or {}
        cal = f.get("calibration") or {}
        vif = (asm.get("collinearity") or {}).get("vif") or {}
        bits = [f"{asm['events_per_parameter']} events per predictor parameter" if asm.get("events_per_parameter") is not None else "",
                f"C-statistic {_x(f.get('c_statistic'))}" if f.get("c_statistic") is not None else "",
                f"calibration slope {_x(cal.get('slope'))} (apparent)" if cal.get("slope") is not None else "",
                f"Hosmer-Lemeshow P = {_p((cal.get('hosmer_lemeshow') or {}).get('p_value'))}"
                if (cal.get("hosmer_lemeshow") or {}).get("p_value") is not None else "",
                f"adjusted R² {_x(f.get('adjusted_r2'))}" if f.get("adjusted_r2") is not None else "",
                f"largest variance inflation factor {_x(max(vif.values()))}" if vif else ""]
        notes.append("Model diagnostics: " + "; ".join(b for b in bits if b) + ".")
    notes.append("; ".join(f"{k}, {v}" for k, v in sorted(abbr.items())) + ".")
    title = str(a["spec"].get("title") or aid)
    return {"analysis_id": aid, "kind": "results", "title": _title_from_request(title) if is_instruction(title) else title,
            "columns": cols, "rows": rows, "footnotes": notes}


def build(s: Any, blobs: BlobStore, with_duplicates: bool = False) -> list[dict[str, Any]]:
    """Every table, numbered in reporting order, from the current (not replaced) analyses. A table identical to an
    earlier one (same headings and cells) is not numbered again; with_duplicates=True returns it, marked."""
    from .workflow import current
    labels = {}
    if s.curator_reports:
        labels = _labels(blobs.get_json(s.curator_reports[-1]["report_sha256"]).get("dictionary"))
    tables = []
    for aid, a in sorted(current(s).items(), key=lambda kv: position(kv[1])):
        if a.get("standard") == "data_quality" or not blobs.has(a["outputs_sha256"]):
            continue
        out = blobs.get_json(a["outputs_sha256"])
        t = _table1(aid, a, out, labels) if a.get("standard") == "table1" else _results(aid, a, out, labels)
        heading = ((a.get("description") or {}).get("heading") or "").strip()
        if t and heading and not a.get("standard") and not t.get("laid_out_by"):
            t["title"] = heading
        elif t and heading and t.get("laid_out_by") and (a.get("basis") == "post hoc" or is_instruction(t["title"])):
            t["title"] = heading           # an author's request is not a table title
        elif t and t.get("laid_out_by") and is_instruction(t["title"]):
            t["title"] = _title_from_request(t["title"])
        if t:
            tables.append(t | {"label": a.get("label")})
    seen: dict[str, dict[str, Any]] = {}
    n = 0
    for t in tables:
        key = json.dumps([t["columns"], [r["cells"] for r in t["rows"]]])
        if key in seen:
            t["number"], t["duplicate_of"] = None, {"analysis_id": seen[key]["analysis_id"], "number": seen[key]["number"]}
        else:
            n += 1
            t["number"] = n
            seen[key] = t
    return tables if with_duplicates else [t for t in tables if t["number"]]


def merge_plan(s: Any, blobs: BlobStore, target: str | None, known: set[str] | None = None) -> dict[str, Any] | None:
    """Folding a table that describes the participants by group into Table 1: Table 1's grouping, its
    characteristics plus those the other table adds, and whether the other table can leave the paper."""
    from .workflow import current
    cur = current(s)
    cands = [t for t in build(s, blobs, with_duplicates=True)
             if t.get("grouped") and t["analysis_id"] in cur and not cur[t["analysis_id"]].get("standard")]
    t = next((x for x in cands if x["analysis_id"] == target), None) or (cands[0] if len(cands) == 1 else None)
    if t is None:
        return None
    if known is None and s.curator_reports:
        known = {d.get("column") for d in blobs.get_json(s.curator_reports[-1]["report_sha256"]).get("dictionary") or []}
    t1 = next((v for v in cur.values() if v.get("standard") == "table1"), None)
    spec1 = dict((t1 or {}).get("spec", {}).get("table1") or (s.plan_locked or {}).get("table1") or {})
    have = list(spec1.get("variables") or [])
    add = [c for c in t["grouped"]["columns"] if (not known or c in known) and c not in have and c != spec1.get("stratify_by")]
    return {"analysis_id": t["analysis_id"], "table_number": t.get("number"), "stratify_by": spec1.get("stratify_by"),
            "variables": have + add, "added": add, "leave_out": cur[t["analysis_id"]].get("basis") == "post hoc"}


def listing(s: Any, blobs: BlobStore) -> str:
    rows = [f"Table {t['number']} [A:{t['analysis_id']}]: {t['title']}" for t in build(s, blobs)]
    return "\n".join(rows) or "(none)"


def to_docx(tables: list[dict[str, Any]], heading: str | None = None) -> bytes:
    """A Word file with the tables in journal style: three rules, no vertical lines, footnotes below."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    def borders(cell, **edges):
        tcpr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for edge, on in edges.items():
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single" if on else "nil")
            if on:
                e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
            b.append(e)
        tcpr.append(b)

    def tight(para):
        f = para.paragraph_format
        f.space_before, f.space_after, f.line_spacing = Pt(1), Pt(1), 1.0

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    if heading:
        doc.add_paragraph().add_run(heading).bold = True
    for t in tables:
        p = doc.add_paragraph()
        p.add_run(f"Table {t['number']}. ").bold = True
        p.add_run(t["title"])
        tb = doc.add_table(rows=1 + len(t["rows"]), cols=len(t["columns"]))
        last = len(t["rows"])
        # Proportionate widths: the label column gets room for its words, the value columns share the rest.
        from docx.shared import Cm
        ncol = len(t["columns"])
        total = 16.0
        first = total * (0.36 if ncol <= 4 else 0.28 if ncol <= 6 else 0.22)
        rest = (total - first) / max(ncol - 1, 1)
        tb.autofit = False
        for row in tb.rows:
            for j, cell in enumerate(row.cells):
                cell.width = Cm(first if j == 0 else rest)
        size = Pt(10 if ncol <= 5 else 9 if ncol <= 7 else 8)
        for j, h in enumerate(t["columns"]):
            c = tb.rows[0].cells[j]
            c.text = ""
            for k, part in enumerate(h.split("\n")):
                run = c.paragraphs[0].add_run(part)
                run.bold, run.font.size = True, size
                if k < len(h.split("\n")) - 1:
                    run.add_break()
            tight(c.paragraphs[0])
            if j:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            borders(c, top=True, bottom=True, left=False, right=False)
        for i, r in enumerate(t["rows"], start=1):
            for j, v in enumerate(r["cells"]):
                c = tb.rows[i].cells[j]
                c.text = ""
                para = c.paragraphs[0]
                para.add_run(("    " if r.get("indent") and j == 0 else "") + str(v)).font.size = size
                mk = (r.get("marks") or {}).get(j) or (r.get("marks") or {}).get(str(j))
                if mk:
                    para.add_run(mk).font.superscript = True
                if j:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                tight(para)
                borders(c, top=False, bottom=(i == last), left=False, right=False)
        for n in t["footnotes"]:
            fp = doc.add_paragraph()
            m = re.match(r"^([a-j]) (.*)$", n)
            if m:
                sup = fp.add_run(m.group(1))
                sup.font.superscript, sup.font.size = True, Pt(8.5)
                n = " " + m.group(2)
            fp.add_run(n).font.size = Pt(8.5)
            fp.paragraph_format.space_after = Pt(0)
        doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# -- Authors' wording on the tables -------------------------------------------------
# Authors can reword a table on the draft page: its title, column headings, row labels and footnotes.
# Numbers come from the results and cannot change. Codes in labels can (a "0" group can become "No GDM").
_NUMS = re.compile(r"-?\d+(?:\.\d+)?")
_N_IN_HEADER = re.compile(r"\(n\s*=\s*(\d+)\)")


def _numbers(text: str) -> list[str]:
    return sorted(_NUMS.findall(text or ""))


def check_edits(tables: list[dict[str, Any]], edits: dict[str, Any]) -> dict[str, Any]:
    """Validate the authors' wording changes against the computed tables; return them tidied.
    Raises ValueError with a plain message when a change would alter a number."""
    by_id = {t["analysis_id"]: t for t in tables}
    clean: dict[str, Any] = {}
    for aid, e in (edits or {}).items():
        t = by_id.get(aid)
        if t is None or not isinstance(e, dict):
            raise ValueError(f"there is no table for {aid!r}")
        out: dict[str, Any] = {}
        label = f"Table {t['number']}"

        def keep(old: str, new: str, where: str, required: list[str] | None = None) -> str:
            new = str(new).strip()
            if not new:
                raise ValueError(f"{label}: {where} cannot be empty")
            need = _numbers(old) if required is None else sorted(required)
            have = _numbers(new)
            missing = [n for n in need if n not in have or have.count(n) < need.count(n)]
            if missing or (required is None and have != need):
                raise ValueError(f"{label}: {where} must keep its numbers ({', '.join(need) or 'none'}); "
                                 "numbers come from the results and are changed only by redoing the analysis")
            return new
        if "title" in e:
            out["title"] = keep(t["title"], e["title"], "the title")
        cols = {}
        for j, v in (e.get("columns") or {}).items():
            j = int(j)
            if not 0 <= j < len(t["columns"]):
                raise ValueError(f"{label}: no column {j}")
            old = t["columns"][j]
            cols[str(j)] = keep(old, v, f"heading {j + 1}", _N_IN_HEADER.findall(old)) if j else str(v).strip() or old
        if cols:
            out["columns"] = cols
        rows = {}
        for i, v in (e.get("rows") or {}).items():
            i = int(i)
            if not 0 <= i < len(t["rows"]):
                raise ValueError(f"{label}: no row {i + 1}")
            new = str(v).strip()
            if not new:
                raise ValueError(f"{label}: row label {i + 1} cannot be empty")
            rows[str(i)] = new                                  # row labels are words (codes may change)
        if rows:
            out["rows"] = rows
        notes = {}
        for k, v in (e.get("footnotes") or {}).items():
            k = int(k)
            if not 0 <= k < len(t["footnotes"]):
                raise ValueError(f"{label}: no footnote {k + 1}")
            notes[str(k)] = keep(t["footnotes"][k], v, f"footnote {k + 1}")
        if notes:
            out["footnotes"] = notes
        if out:
            clean[aid] = out
    return clean


def apply_edits(tables: list[dict[str, Any]], edits: dict[str, Any] | None) -> list[dict[str, Any]]:
    """The tables with the authors' wording, each changed piece marked so the page can show it."""
    out = []
    for t in tables:
        e = (edits or {}).get(t["analysis_id"]) or {}
        t = {**t, "columns": list(t["columns"]), "rows": [dict(r, cells=list(r["cells"])) for r in t["rows"]],
             "footnotes": list(t["footnotes"]), "edited": sorted(k for k in e)}
        if e.get("title"):
            t["title"] = e["title"]
        for j, v in (e.get("columns") or {}).items():
            if int(j) < len(t["columns"]):
                t["columns"][int(j)] = v
        for i, v in (e.get("rows") or {}).items():
            if int(i) < len(t["rows"]):
                t["rows"][int(i)]["cells"][0] = v
        for k, v in (e.get("footnotes") or {}).items():
            if int(k) < len(t["footnotes"]):
                t["footnotes"][int(k)] = v
        out.append(t)
    return out


def resolve(text: str, aid: str, tables: list[dict[str, Any]], figures: list[dict[str, Any]]) -> str:
    """Replace a description's [[Table]] and [[Figure:name]] tokens with the current numbers."""
    t = next((x for x in tables if x["analysis_id"] == aid), None)
    out = (text or "").replace("[[Table]]", f"Table {t['number']}" if t else "the table")
    for f in figures:
        if f["analysis_id"] == aid:
            out = out.replace(f"[[Figure:{f['name']}]]", f"Figure {f['number']}" if f.get("number") else "the figure")
    return re.sub(r"\[\[Figure:[^\]]*\]\]", "the figure", out)


_REF = re.compile(r"\b(Tables?|Fig(?:ure)?s?\.?)\s+(\d+(?:\s*(?:,|and|to|-|–)\s*\d+)*)", re.I)


def first_mentions(text: str) -> list[tuple[str, int]]:
    """Tables and figures in the order the text first refers to them."""
    seen: list[tuple[str, int]] = []
    for m in _REF.finditer(text or ""):
        kind = "table" if m.group(1).lower().startswith("tab") else "figure"
        nums = [int(x) for x in re.findall(r"\d+", m.group(2))]
        for a, b in re.findall(r"(\d+)\s*(?:to|-|–)\s*(\d+)", m.group(2)):
            nums += list(range(int(a), int(b) + 1))
        for n in sorted(set(nums)):
            if (kind, n) not in seen:
                seen.append((kind, n))
    return seen


def out_of_order(text: str) -> list[str]:
    """Tables or figures first referred to before a lower-numbered one (journals number them in order of mention)."""
    bad, top = [], {"table": 0, "figure": 0}
    for kind, n in first_mentions(text):
        if n < top[kind]:
            bad.append(f"{kind.capitalize()} {n}")
        top[kind] = max(top[kind], n)
    return bad


def references(text: str) -> dict[str, set[int]]:
    """Table and figure numbers the text refers to, ranges expanded ("Tables 2 to 4", "Figures 1 and 2")."""
    found: dict[str, set[int]] = {"table": set(), "figure": set()}
    for m in _REF.finditer(text or ""):
        kind = "table" if m.group(1).lower().startswith("tab") else "figure"
        nums = [int(x) for x in re.findall(r"\d+", m.group(2))]
        ranges = re.findall(r"(\d+)\s*(?:to|-|–)\s*(\d+)", m.group(2))
        found[kind].update(nums)
        for a, b in ranges:
            found[kind].update(range(int(a), int(b) + 1))
    return found
