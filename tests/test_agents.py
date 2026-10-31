# SPDX-License-Identifier: AGPL-3.0-only
# Copyright (c) 2026 Boon-How Chew

"""Agents end to end with a scripted model and the local executor."""

import json
from pathlib import Path

import re

import pytest

from saksi import demo_data
from saksi.agents import AgentError, Agents
from saksi.blobs import BlobStore
from saksi.canonical import sha256_bytes
from saksi.llm import ModelClient
from saksi.sandbox import LocalExecutor
from saksi.workflow import WorkflowError, ledger
from tests.test_sandbox import GOOD

AUC_BROKEN = '''
import pandas as pd
def analyse(data_path, seed):
    df = pd.read_csv(data_path)
    return {"x": float(df["no_such_column"].mean())}
'''
AUC_FIXED = '''
import numpy as np
import pandas as pd
import statsmodels.formula.api as smf
from sklearn.metrics import roc_auc_score

def analyse(data_path, seed):
    df = pd.read_csv(data_path).dropna(subset=["gdm", "booking_bmi", "age", "parity", "previous_gdm"])
    m = smf.logit("gdm ~ booking_bmi + age + parity + previous_gdm", data=df).fit(disp=0)
    p = m.predict(df)
    rng = np.random.default_rng(seed)
    boots = []
    y = df["gdm"].to_numpy(); pv = p.to_numpy()
    for _ in range(200):
        i = rng.integers(0, len(y), len(y))
        if len(set(y[i])) == 2:
            boots.append(roc_auc_score(y[i], pv[i]))
    return {"n_analysed": int(len(df)), "n_excluded": 0, "exclusion_reason": "",
            "estimates": [{"name": "AUC", "measure": "AUC", "estimate": float(roc_auc_score(y, pv)),
                           "ci_lower": float(np.percentile(boots, 2.5)), "ci_upper": float(np.percentile(boots, 97.5)),
                           "p_value": None}],
            "model": {"formula_or_description": "same as primary", "converged": True}, "notes": []}


def figures(data_path, seed):
    import matplotlib.pyplot as plt
    from sklearn.metrics import roc_curve
    df = pd.read_csv(data_path).dropna(subset=["gdm", "booking_bmi", "age", "parity", "previous_gdm"])
    m = smf.logit("gdm ~ booking_bmi + age + parity + previous_gdm", data=df).fit(disp=0)
    fpr, tpr, _ = roc_curve(df["gdm"], m.predict(df))
    fig, ax = plt.subplots(figsize=(4, 4))
    ax.plot(fpr, tpr); ax.plot([0, 1], [0, 1], linestyle="--")
    ax.set_xlabel("1 - specificity"); ax.set_ylabel("Sensitivity")
    return {"roc_curve": {"figure": fig, "caption": "ROC curve of the primary model in 388 women."}}
'''
CURATOR_REPLY = {
    "summary": "400 records, 10 columns. Two direct identifiers should be removed.",
    "dictionary": [{"column": "gdm", "label": "Gestational diabetes", "type": "binary", "units_or_coding": "0/1",
                    "role": "outcome"},
                   {"column": "booking_bmi", "label": "BMI at booking", "type": "numeric", "units_or_coding": "kg/m2",
                    "role": "exposure"}],
    "flags": [{"column": "patient_name", "severity": "high", "issue": "direct identifier",
               "suggested_check": "remove before analysis"}],
    "protocol_variables_missing_from_data": [],
}
PLAN_REPLY = {"analyses": [
    {"id": "primary", "title": "BMI and GDM", "objective": "association", "outcome": {"name": "GDM", "column": "gdm"},
     "exposures": [{"name": "BMI", "column": "booking_bmi"}],
     "covariates": [{"name": "age", "column": "age"}, {"name": "parity", "column": "parity"},
                    {"name": "previous GDM", "column": "previous_gdm"}],
     "method": "multivariable logistic regression", "effect_measure": "adjusted odds ratio",
     "missing_data": "complete case",
     "kind": "other",     # these older tests do not ask for diagnostics; the standard-sequence tests set "multivariable"
     "protocol_quote": "Multivariable logistic regression with GDM as the outcome and booking BMI as the exposure",
     "role": "primary",
     "primary_quote": "Primary objective: To estimate the association between booking body mass index (BMI) and "
                      "gestational diabetes mellitus (GDM)",
     "sample_size_quote": "400 women (about 40 GDM events) give 80% power to detect an odds ratio of 1.10 per kg/m2",
     "sample_size_target": {"n": 400, "events": 40},
     "unresolved": []},
    {"id": "auc", "title": "Discrimination", "objective": "AUC", "outcome": {"name": "GDM", "column": "gdm"},
     "exposures": [], "covariates": [{"name": "waist", "column": "waist_cm"}], "method": "AUC with 95% CI",
     "effect_measure": "AUC", "missing_data": "complete case",
     "protocol_quote": "The model was externally validated in three cohorts", "role": "secondary",
     "primary_quote": "", "sample_size_quote": "", "sample_size_target": {"n": None, "events": None},
     "unresolved": []},
]}


class ScriptedModel:
    """Answers by task, in order. Records every request body it receives."""

    def __init__(self, script):
        self.script = {k: list(v) for k, v in script.items()}
        self.bodies = []

    def post(self, url, headers, json, timeout):
        self.bodies.append(json)
        system = json["messages"][0]["content"]
        if "were not found in the protocol" in system:
            r = (self.script.get("quotes") or ['{"fixes": []}']).pop(0)
            r = r if isinstance(r, str) else __import__("json").dumps(r)
            return {"choices": [{"message": {"content": r}, "finish_reason": "stop"}],
                    "usage": {"prompt_tokens": 10, "completion_tokens": 5}}
        purpose = ("review" if "Reviewer Assistant" in system else "curate" if "Data Curator" in system
                   else "plan" if "structured analysis plan" in system
                   else "describe" if "reports ONE analysis" in system
                   else "assist" if "Saksi's analysis assistant" in system
                   else "plan_assist" if "settle what this particular paper reports" in system
                   else "checklist" if "against a reporting guideline" in system
                   else "write" if "Writing Agent" in system else "code")
        if purpose == "describe" and not self.script.get("describe"):
            reply = describe_reply                                   # a plain, number-free default
        elif purpose == "assist" and not self.script.get("assist"):
            reply = assist_reply
        elif purpose == "plan_assist" and not self.script.get("plan_assist"):
            reply = {"reply": "Then this paper reports the primary objective only.", "questions": ["Is GDM defined by the 2013 WHO criteria?"],
                     "redraft": "This paper reports the primary objective only."}
        elif purpose == "checklist" and not self.script.get("checklist"):
            reply = checklist_reply
        else:
            reply = self.script[purpose].pop(0)
        if callable(reply):
            reply = reply(json["messages"][1]["content"])
        reply, finish = reply if isinstance(reply, tuple) else (reply, "stop")     # (text, "length") = cut off
        content = reply if isinstance(reply, str) else __import__("json").dumps(reply)
        return {"choices": [{"message": {"content": content}, "finish_reason": finish}],
                "usage": {"prompt_tokens": 100, "completion_tokens": 50}}


def checklist_reply(user):
    """Scripted checklist: every item 'reported' with the draft's first words, except one 'missing'."""
    import json as _j
    items = _j.loads(user.split("ITEMS:\n", 1)[1].split("\n\nDRAFT:", 1)[0])
    draft = user.split("DRAFT:\n", 1)[1]
    first = next(ln for ln in draft.splitlines() if ln and not ln.startswith("#") and not ln.startswith("["))[:30]
    return _j.dumps({"items": [{"id": it["id"], "status": "missing" if i == 0 else "reported",
                                "quote": "" if i == 0 else (first if i % 2 else "words not in the draft at all"),
                                "note": "[Authors] the setting"} for i, it in enumerate(items)]})


def assist_reply(user):
    """The assistant, scripted: 'redo' -> redo, 'text' -> rewrite_text, otherwise a variant."""
    import json as _j
    msg = user.split("AUTHOR'S MESSAGE: ", 1)[1].split("\n\nABOUT: ", 1)[0]
    about = user.split("ABOUT: ", 1)[1].split("\n", 1)[0]
    aid = None if about.startswith("the Results") else about
    low = msg.lower()
    kind = "redo" if "redo" in low else "rewrite_text" if "text" in low else "variant"
    return _j.dumps({"reply": f"I will {'redo' if kind == 'redo' else 'rewrite the text of' if kind == 'rewrite_text' else 'run a variant of'} it.",
                     "action": {"type": kind, "analysis_id": aid, "instruction": msg}})


def describe_reply(user):
    """A Results (or Methods) text for any analysis, from what the describe request gives the model."""
    import json as _j, re as _re
    info = _j.loads(user.split("ANALYSIS:\n", 1)[1].split("\n\nRESULTS:", 1)[0])
    out = _j.loads(user.split("RESULTS:\n", 1)[1].split("\n\nTABLE:", 1)[0])
    figs = _j.loads(user.split("FIGURES:\n", 1)[1].split("\n\nINTERVAL FACTS", 1)[0])
    aid, refs = info["id"], " and ".join(["[[Table]]"] + [f["token"] for f in figs])
    if info.get("standard") == "data_quality":
        miss = [c for c in out.get("columns", []) if c.get("missing")]
        return (f"Before analysis, the {out['n_rows']} records were checked for missing values, duplicate records and "
                f"outliers beyond 1.5 times the interquartile range, which were counted and not removed [A:{aid}]. "
                f"There were {out['duplicate_rows']} duplicated records"
                + (f", and {miss[0]['column']} was missing for {miss[0]['missing']}" if miss else "") + f" [A:{aid}].")
    if info.get("standard") == "table1":
        return (f"Heading: Characteristics of the participants\n\n[[Table]] shows the characteristics of the {out['n_analysed']} participants by "
                f"{out.get('stratify_by') or 'group'} [A:{aid}].")
    est = (out.get("estimates") or [None])[0]
    lead = "In this exploratory analysis, " if info["label"] == "exploratory" else ""
    if not est:
        return f"{lead}{refs} report this analysis [A:{aid}]."
    ci = (f" (95% CI {est['ci_lower']:.2f} to {est['ci_upper']:.2f})" if est.get("ci_lower") is not None else "")
    what = est["measure"] if est["name"] == est["measure"] else f"{est['measure']} for {est['name']}"
    sentence = f"{lead}the {what} was {est['estimate']:.2f}{ci} [A:{aid}]."
    return f"Heading: {info['title'][:60]}\n\n{refs} give the results. {sentence[0].upper()}{sentence[1:]}"


def fence(code):
    return f"```python\n{code}```"


def upload(p, kind, filename, data):
    sha = sha256_bytes(data)
    dest = p.root / "materials" / sha
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(data)
    p.upload_material("bhc", kind, filename, sha, len(data), storage_ref=f"materials/{sha}")
    return sha


@pytest.fixture
def setup(proj, monkeypatch, tmp_path):
    monkeypatch.setenv("META_API_KEY", "test-key")
    model = ScriptedModel({
        "curate": [CURATOR_REPLY],
        "plan": [PLAN_REPLY],
        "code": [fence(GOOD), fence(AUC_BROKEN), fence(AUC_FIXED), fence(GOOD)],
    })
    agents = Agents(proj, ModelClient(proj.blobs, transport=model, mode="live"), LocalExecutor())
    data = demo_data.make().to_csv(index=False).encode()
    ds_sha = upload(proj, "dataset", "gdm_synthetic.csv", data)
    upload(proj, "protocol", "protocol.txt", demo_data.PROTOCOL.encode())
    return proj, agents, model, ds_sha


def test_curate_plan_run_end_to_end(setup):
    p, agents, model, ds_sha = setup

    s = agents.curate("bhc")
    rep = p.blobs.get_json(s.curator_reports[-1]["report_sha256"])
    assert rep["flags"][0]["column"] == "patient_name" and rep["prompt"]["name"] == "curator"
    assert s.curator_reports[-1]["material_sha256"] == ds_sha

    s = agents.propose_plan("coauthor")
    plan = s.plan["analyses"]
    assert plan["primary"]["protocol_quote_verified"] is True
    assert plan["auc"]["protocol_quote_verified"] is False        # invented quote is caught
    assert plan["auc"]["columns_not_in_data"] == ["waist_cm"]     # invented column is caught

    for a in ("bhc", "coauthor"):
        s = p.confirm_plan(a, s.plan["sha"])
    assert s.stage == "ANALYSIS"

    s = agents.run_planned("bhc")
    prim, auc = s.analyses["primary"], s.analyses["auc"]
    # Planned as the confirmatory primary analysis, but the achieved sample fell short: exploratory (author's
    # definition, 31 Oct), decided from the counts alone.
    assert prim["status"] == "ok" and prim["label"] == "exploratory" and prim["basis"] == "pre-specified"
    assert "fell short" in prim["downgraded"]
    # Planned for 400 women and 40 events; 388 and 18 were analysed. Flagged; the label stays.
    assert prim["sample_check"]["shortfall"] is True
    assert prim["sample_check"]["targets"] == {"n": {"target": 400, "achieved": 388, "short": True},
                                               "events": {"target": 40, "achieved": 18, "short": True}}
    assert auc["label"] == "exploratory" and auc["basis"] == "pre-specified"     # secondary: not powered for it
    assert any("listed as a secondary analysis" in r
               for r in s.plan_locked["analyses"]["auc"]["design"]["reasons_not_confirmatory"])

    # Figures: Saksi draws the forest plot and flow diagram from the recorded results; the ROC curve,
    # required by rule for an AUC analysis, is drawn by the analysis code in the sandbox.
    figs = {(f["analysis_id"], f["name"]): f for f in s.figures}
    assert figs[("primary", "forest_plot")]["drawn_by"] == "saksi" and figs[("primary", "forest_plot")]["number"] == 1
    assert figs[("primary", "participant_flow")]["drawn_by"] == "saksi"
    assert auc["spec"]["required_figure"]["name"] == "roc_curve" and auc["execution"]["required_figure_drawn"]
    roc = figs[("auc", "roc_curve")]
    assert roc["drawn_by"] == "analysis code" and roc["label"] == "exploratory"
    # The code caption is kept as written; the label belongs in the text.
    assert roc["caption"].startswith("ROC curve")              # no label in captions (author, 31 Oct)
    assert roc["caption_from_code"] == "ROC curve of the primary model in 388 women."
    assert roc["caption_check"]["ok"] is True
    assert p.blobs.get_text(roc["svg_sha256"]).lstrip().startswith("<?xml")
    assert all(f["caption_check"]["untraced"] == 0 for f in s.figures)
    # AUC_FIXED claims 0 excluded while 12 of 400 records were dropped: Saksi flags the flow diagram.
    assert "do not add up" in figs[("auc", "participant_flow")]["warning"]
    assert figs[("primary", "participant_flow")]["warning"] is None
    assert len(prim["execution"]["attempts"]) == 1
    assert prim["execution"]["provenance"]["data_sha256"] == ds_sha
    assert auc["status"] == "ok" and [a["outcome"] for a in auc["execution"]["attempts"]] == ["error", "ok"]
    assert len(auc["llm_call_seqs"]) == 2                           # failed attempt and fix both logged
    out = p.blobs.get_json(auc["outputs_sha256"])
    assert 0.5 < out["estimates"][0]["estimate"] < 1.0

    s = agents.run_request("coauthor", "q1", "Same model stratified by ethnicity")
    x = s.analyses["x_q1"]
    assert x["label"] == "exploratory" and x["basis"] == "post hoc" and x["requested_by"] == "coauthor"

    for a in s.authors:
        s = p.accept_analyses(a, s.analysis_set_sha())
    assert s.stage == "DRAFTING"
    assert p.log.verify().ok

    led = ledger(p.log.events())
    assert led["agent:analyst"]["on_behalf_of"] == {"coauthor": 2, "bhc": 2}
    assert led["agent:analyst"]["counts"]["model_calls"] == 6     # plan + quote repair + primary + auc x2 + request

    # The models never received a data row.
    sent = json.dumps(model.bodies)
    for leak in ("Person 1", "S100001", "S100399"):
        assert leak not in sent


def test_rejected_code_is_recorded_as_failure(setup):
    p, agents, model, _ = setup
    model.script["code"] = [fence("import os\ndef analyse(data_path, seed):\n    return {}")] * 2
    agents.curate("bhc")
    model.script["plan"] = [{"analyses": [PLAN_REPLY["analyses"][0]]}]
    s = agents.propose_plan("bhc")
    for a in ("bhc", "coauthor"):
        s = p.confirm_plan(a, s.plan["sha"])
    s = agents.run_planned("bhc")
    a = s.analyses["primary"]
    assert a["status"] == "error" and "import of 'os'" in a["error"]
    assert [x["outcome"] for x in a["execution"]["attempts"]] == ["rejected", "rejected"]
    # A rerun gets a new id; the failure stays in the record.
    model.script["code"] = [fence(GOOD)]
    s = agents.run_planned("bhc")
    assert s.analyses["primary.r1"]["status"] == "ok" and s.analyses["primary"]["status"] == "error"


def test_bad_json_from_curator_is_retried_once_then_refused(setup):
    p, agents, model, _ = setup
    model.script["curate"] = ["sorry, I cannot", "still not JSON"]
    with pytest.raises(AgentError, match="after one retry"):
        agents.curate("bhc")
    assert [e.type for e in p.log.events()[-2:]] == ["llm.called", "llm.called"]
    assert "YOUR PREVIOUS REPLY WAS NOT VALID JSON" in model.bodies[-1]["messages"][1]["content"]


def test_curator_retry_can_succeed(setup):
    p, agents, model, _ = setup
    model.script["curate"] = ['{"summary": "x", "dictionary": [', CURATOR_REPLY]
    s = agents.curate("bhc")
    assert s.curator_reports[-1]["summary"].startswith("400 records")
    assert len(p.blobs.get_json(s.curator_reports[-1]["report_sha256"])["earlier_call_seqs"]) == 1


def test_wide_dataset_is_curated_in_batches_and_a_cut_off_batch_is_split(proj, monkeypatch):
    """29 Oct, live: a real dataset with many columns broke the Data Curator. Its single reply was cut
    off at the length limit mid-dictionary ("Expecting ',' delimiter")."""
    import numpy as np
    import pandas as pd
    monkeypatch.setenv("META_API_KEY", "k")
    rng = np.random.default_rng(1)
    wide = pd.DataFrame({f"v{i:03d}": rng.normal(size=60).round(2) for i in range(90)})
    upload(proj, "dataset", "wide.csv", wide.to_csv(index=False).encode())
    upload(proj, "protocol", "protocol.txt", demo_data.PROTOCOL.encode())
    entry = lambda a, b: {"dictionary": [{"column": f"v{i:03d}", "label": "x", "type": "numeric",
                                          "units_or_coding": "unknown", "role": "unclear"} for i in range(a, b)],
                          "flags": [{"column": f"v{a:03d}", "severity": "low", "issue": "check", "suggested_check": "c"}]}
    model = ScriptedModel({"curate": [
        ('{"dictionary": [{"column": "v000"}', "length"),          # first 40 columns: cut off
        entry(0, 20), entry(20, 40), entry(40, 80), entry(80, 90),
        {"summary": "90 numeric columns.", "protocol_variables_missing_from_data": ["gdm"], "flags": []}]})
    agents = Agents(proj, ModelClient(proj.blobs, transport=model, mode="live"), LocalExecutor())
    s = agents.curate("bhc")
    rep = proj.blobs.get_json(s.curator_reports[-1]["report_sha256"])
    assert rep["columns"] == 90 and len(rep["dictionary"]) == 90 and len(rep["flags"]) == 4
    assert rep["summary"] == "90 numeric columns." and rep["protocol_variables_missing_from_data"] == ["gdm"]
    assert [x["name"] for x in rep["prompts"]] == ["curator_columns", "curator_overview"]
    calls = [e for e in proj.log.events() if e.type == "llm.called"]
    assert len(calls) == 6                                      # the cut-off reply is in the record too
    assert "cut off" not in json.dumps(model.bodies[-1]) and "v089" in model.bodies[-1]["messages"][1]["content"]


def test_tampered_material_is_refused(setup):
    p, agents, model, ds_sha = setup
    path = p.root / "materials" / ds_sha
    path.write_bytes(path.read_bytes() + b"\n1,2,3")
    with pytest.raises(AgentError, match="does not match"):
        agents.curate("bhc")


def test_returned_error_is_labelled_honestly(setup):
    p, agents, model, _ = setup
    cannot = "def analyse(data_path, seed):\n    return {'error': 'column waist_cm not found'}\n"
    model.script["code"] = [fence(cannot), fence(GOOD)]
    agents.curate("bhc")
    model.script["plan"] = [{"analyses": [PLAN_REPLY["analyses"][0]]}]
    s = agents.propose_plan("bhc")
    for a in ("bhc", "coauthor"):
        s = p.confirm_plan(a, s.plan["sha"])
    s = agents.run_planned("bhc")
    a = s.analyses["primary"]
    assert a["status"] == "ok"
    assert [x["outcome"] for x in a["execution"]["attempts"]] == ["returned_error", "ok"]


def _to_drafting(p, agents, model):
    agents.curate("bhc")
    model.script["plan"] = [{"analyses": [PLAN_REPLY["analyses"][0]]}]
    s = agents.propose_plan("bhc")
    for a in ("bhc", "coauthor"):
        s = p.confirm_plan(a, s.plan["sha"])
    model.script["code"] = [fence(GOOD), fence(GOOD)]
    agents.run_planned("bhc")
    s = agents.run_request("coauthor", "q1", "Same model, any extra")
    for a in s.authors:
        s = p.accept_analyses(a, s.analysis_set_sha())
    assert s.stage == "DRAFTING"
    est = {e["name"]: e for e in p.blobs.get_json(s.analyses["primary"]["outputs_sha256"])["estimates"]}
    # Every table and figure must be referred to in the text.
    from saksi import tables
    nt, nf = len(tables.build(s, p.blobs)), max(f["number"] or 0 for f in s.figures)
    global REFS
    REFS = f"Tables 1 to {nt} and Figures 1 to {nf} give the estimates. "
    return est["booking_bmi"]


REFS = ""


def _draft(bmi, extra="The adjusted OR for previous GDM was 9.99 [A:primary]. ",
           exploratory="A further model gave a similar estimate [A:x_q1]."):
    return ("[Authors: introduction]\n\n## Methods\nA multivariable logistic regression was pre-specified.\n\n"
            "## Results\n"
            f"Of 400 records, none was duplicated and 12 lacked booking BMI [A:s_data_quality]. "
            f"Of 400 records, 388 were analysed, with 18 GDM events against 40 planned [A:primary]. "
            f"Booking BMI was associated with GDM "
            f"(adjusted OR {bmi['estimate']:.2f}, 95% CI {bmi['ci_lower']:.2f} to {bmi['ci_upper']:.2f}) [A:primary]. "
            f"{REFS}{extra}{exploratory}\n\n## Discussion\nThe estimate was imprecise.\n")


def test_writer_fixes_its_own_draft_after_check(setup):
    p, agents, model, _ = setup
    bmi = _to_drafting(p, agents, model)
    model.script["write"] = [_draft(bmi),
                             _draft(bmi, extra="", exploratory="An exploratory model gave a similar estimate [A:x_q1].")]
    s = agents.write_draft("bhc")
    d = s.drafts[1]
    assert d["by"] == "agent:writer" and d["requested_by"] == "bhc"
    assert d["sections"] == ["Methods", "Results", "Discussion"]
    assert d["checks"]["ok"] is True
    assert [a["passed_check"] for a in d["writer_attempts"]] == [False, True]
    first = p.blobs.get_json(d["writer_attempts"][0]["findings_sha256"])
    assert [u["number"] for u in first["untraced"]] == ["9.99"] and first["exploratory_not_labelled"]
    assert d["writer_attempts"][1]["prompt"]["name"] == "rewrite"
    # The rewrite request told the agent exactly what was wrong.
    assert '"interval_includes_null": true' in model.bodies[-2]["messages"][1]["content"]   # facts, not judgement
    assert "Figure 1 [A:primary]: Adjusted odds ratio with 95% confidence intervals" in model.bodies[-2]["messages"][1]["content"]
    rewrite_msg = model.bodies[-1]["messages"][1]["content"]
    assert "UNTRACED 9.99" in rewrite_msg and "EXPLORATORY NOT LABELLED" in rewrite_msg
    # Both model calls are in the log: the failed attempt and the final one.
    calls = [e for e in p.log.events() if e.type == "llm.called" and e.actor_id == "agent:writer"]
    assert len(calls) == 2


def test_writer_cannot_pass_by_dropping_an_analysis(setup):
    p, agents, model, _ = setup
    bmi = _to_drafting(p, agents, model)
    # Second attempt "fixes" things by deleting the awkward analysis.
    model.script["write"] = [_draft(bmi), _draft(bmi, extra="", exploratory="")]
    s = agents.write_draft("bhc")
    d = s.drafts[1]
    assert d["checks"]["ok"] is False and d["checks"]["analyses_not_reported"] == ["x_q1"]
    assert [a["passed_check"] for a in d["writer_attempts"]] == [False, False]
    assert s.stage == "AUTHOR_REVISION"          # escalated to the authors, with the findings


def test_unresolved_draft_goes_to_authors_who_can_fix_it(setup):
    from tests.conftest import ANSWERS, RESTATE
    p, agents, model, _ = setup
    bmi = _to_drafting(p, agents, model)
    model.script["write"] = [_draft(bmi), _draft(bmi)]       # agent fails twice
    s = agents.write_draft("bhc")
    d = s.drafts[1]
    assert d["checks"]["untraced"] == 1 and d["checks"]["exploratory_not_labelled"] >= 1
    full = p.blobs.get_json(d["checks_sha256"])
    assert full["untraced"][0]["number"] == "9.99"
    # An author edit that fixes both problems is re-checked automatically.
    text = p.blobs.get_text(d["text_sha256"])
    text = text.replace("The adjusted OR for previous GDM was 9.99 [A:primary]. ", "")
    text = text.replace("A further model", "A further exploratory model")
    s = p.edit_draft("coauthor", 1, text, "removed unsupported figure; labelled exploratory analysis")
    assert s.drafts[2]["checks"]["ok"] is True
    # The sign-off records the check of the exact version signed.
    p.open_draft("bhc", 2)
    s = p.sign_off("bhc", 2, RESTATE, ANSWERS)
    assert s.signoffs[2]["bhc"]["checks_at_signoff"]["ok"] is True


def test_wide_dataset_analysis_sees_only_its_columns():
    from saksi.agents import _focus
    prof = {"rows": 10, "columns": [{"name": f"v{i}"} for i in range(100)]}
    d = {"profile": prof, "dictionary": [{"column": f"v{i}"} for i in range(100)]}
    dic, p, names = _focus(d, [{"outcome": {"column": "v1"}, "exposures": [{"column": "v2"}], "covariates": []}])
    assert [c["name"] for c in p["columns"]] == ["v1", "v2"] and len(dic) == 2 and len(names) == 100
    assert _focus(d, [{"title": "author request"}])[2] is None          # nothing named: everything


def test_plan_cut_off_gives_clear_advice_and_invalid_plan_is_retried(setup):
    p, agents, model, _ = setup
    agents.curate("bhc")
    model.script["plan"] = [('{"analyses": [{"id": "a1"', "length")]
    with pytest.raises(AgentError, match="one objective at a time"):
        agents.propose_plan("bhc")
    model.script["plan"] = ["not json", {"analyses": [PLAN_REPLY["analyses"][0]]}]
    s = agents.propose_plan("bhc", instructions="Only the primary objective")
    assert list(s.plan["analyses"]) == ["primary"] and s.plan["note"].endswith("Only the primary objective")
    assert "THE AUTHORS' INSTRUCTIONS" in model.bodies[-1]["messages"][1]["content"]


def test_unfound_quotations_are_repaired_once_and_rechecked(setup):
    p, agents, model, _ = setup
    agents.curate("bhc")
    bad = dict(PLAN_REPLY["analyses"][0], protocol_quote="Logistic regression of GDM on BMI with the usual adjusters")
    invented = dict(PLAN_REPLY["analyses"][1], role="secondary")         # its quote is invented too
    model.script["plan"] = [{"analyses": [bad, invented]}]
    model.script["quotes"] = [{"fixes": [
        {"id": "primary", "field": "protocol_quote",
         "quote": "Multivariable logistic regression with GDM as the outcome and booking BMI as the exposure"},
        {"id": "auc", "field": "protocol_quote", "quote": "The model was validated in three external cohorts"}]}]
    s = agents.propose_plan("bhc")
    assert s.plan["analyses"]["primary"]["protocol_quote_verified"] is True        # repaired, then re-checked
    assert s.plan["analyses"]["primary"]["label"] == "confirmatory"
    assert s.plan["analyses"]["auc"]["protocol_quote_verified"] is False           # an invented fix is not accepted
    assert "corrected automatically and found: 1" in s.plan["note"]
    assert any(e.payload.get("purpose") == "repair_quotes" for e in p.log.events() if e.type == "llm.called")


def test_prepare_plan_checks_the_data_then_drafts(setup):
    p, agents, model, _ = setup
    model.script["plan"] = [{"analyses": [PLAN_REPLY["analyses"][0]]}]
    s = agents.prepare_plan("bhc")
    assert len(s.curator_reports) == 1 and s.plan and s.plan["requested_by"] == "bhc"
    model.script["plan"] = [{"analyses": [PLAN_REPLY["analyses"][0]]}]
    s = agents.prepare_plan("bhc")                       # already checked: does not curate again
    assert len(s.curator_reports) == 1


def test_wrong_upload_can_be_removed_before_the_plan_is_locked(setup):
    from saksi.workflow import WorkflowError
    p, agents, model, ds_sha = setup
    wrong = upload(p, "dataset", "wrong.csv", b"a,b\n1,2\n")
    assert agents._latest(p.state(), "dataset")["sha256"] == wrong
    with pytest.raises(WorkflowError, match="say why"):
        p.withdraw_material("bhc", wrong, "")
    s = p.withdraw_material("coauthor", wrong, "Uploaded the wrong file")
    assert not (p.root / "materials" / wrong).exists()                       # stored copy deleted
    assert agents._latest(s, "dataset")["sha256"] == ds_sha                  # the right one is in use again
    assert [m["withdrawn"]["reason"] for m in s.materials if m.get("withdrawn")] == ["Uploaded the wrong file"]
    # A plan drafted from a protocol that is then removed cannot be confirmed.
    model.script["plan"] = [{"analyses": [PLAN_REPLY["analyses"][0]]}]
    s = agents.prepare_plan("bhc")
    proto = agents._latest(s, "protocol")["sha256"]
    p.withdraw_material("bhc", proto, "Wrong version of the protocol")
    upload(p, "protocol", "protocol-v2.txt", (demo_data.PROTOCOL + "\nVersion 2.").encode())
    with pytest.raises(WorkflowError, match="protocol that has since been removed"):
        p.confirm_plan("bhc", s.plan["sha"])


def test_standard_sequence_quality_table1_descriptive_then_multivariable(setup):
    """Data quality and Table 1 by Saksi's own code first; then the plan's analyses by kind, whatever the
    protocol order; a multivariable model without diagnostics is sent back once."""
    from tests.test_sandbox import GOOD_WITH_DIAGNOSTICS
    p, agents, model, _ = setup
    agents.curate("bhc")
    multi = dict(PLAN_REPLY["analyses"][0], kind="multivariable")
    descr = {"id": "gdm_by_clinic", "title": "GDM by clinic", "kind": "descriptive", "role": "other",
             "outcome": {"name": "GDM", "column": "gdm"}, "exposures": [{"name": "clinic", "column": "clinic"}],
             "covariates": [], "method": "proportions", "effect_measure": "proportion", "missing_data": "none",
             "protocol_quote": "Multivariable logistic regression with GDM as the outcome", "unresolved": []}
    model.script["plan"] = [{"analyses": [multi, descr],
                             "table1": {"stratify_by": "gdm", "variables": ["age", "parity", "ethnicity", "nope"]}}]
    proportions = ("import pandas as pd\ndef analyse(data_path, seed):\n    df = pd.read_csv(data_path)\n"
                   "    t = df.groupby('clinic')['gdm'].mean()\n"
                   "    return {'n_analysed': int(len(df)), 'estimates': [{'name': str(k), 'measure': 'proportion', "
                   "'estimate': float(v), 'ci_lower': None, 'ci_upper': None, 'p_value': None} for k, v in t.items()]}\n")
    model.script["code"] = [fence(proportions), fence(GOOD), fence(GOOD_WITH_DIAGNOSTICS)]
    s = agents.propose_plan("bhc")
    assert s.plan["table1"]["variables"] == ["age", "parity", "ethnicity"]          # unknown column dropped
    for a in ("bhc", "coauthor"):
        s = p.confirm_plan(a, s.plan["sha"])
    s = agents.run_planned("bhc")
    order = [k for k, _ in sorted(s.analyses.items(), key=lambda kv: kv[1]["seq"])]
    assert order == ["s_data_quality", "s_table1", "gdm_by_clinic", "primary"]
    assert s.analyses["s_table1"]["label"] == "descriptive" and s.analyses["s_table1"]["by"] == "system"
    t1 = p.blobs.get_json(s.analyses["s_table1"]["outputs_sha256"])
    assert [r["variable"] for r in t1["rows"]] == ["age", "parity", "ethnicity"] and t1["groups"][1]["n"] == 20
    prim = s.analyses["primary"]
    assert [x["outcome"] for x in prim["execution"]["attempts"]] == ["missing_diagnostics", "ok"]
    assert "saksi_checks" in model.bodies[-1]["messages"][1]["content"]                # the repair said how
    assert any(f["name"] == "calibration_plot" for f in s.figures if f["analysis_id"] == "primary")
    out = p.blobs.get_json(prim["outputs_sha256"])
    # Diagnostics numbers can be cited and are checked like any other.
    r = p.check_text("There were 4.5 outcome events per predictor parameter [A:primary].")
    assert r["untraced"] == [] and out["diagnostics"]["assumptions"]["events_per_parameter"] == 4.5


def test_diagnostics_still_missing_after_repair_are_flagged_not_hidden(setup):
    p, agents, model, _ = setup
    agents.curate("bhc")
    model.script["plan"] = [{"analyses": [dict(PLAN_REPLY["analyses"][0], kind="multivariable")]}]
    model.script["code"] = [fence(GOOD), fence(GOOD)]
    s = agents.propose_plan("bhc")
    for a in ("bhc", "coauthor"):
        s = p.confirm_plan(a, s.plan["sha"])
    s = agents.run_planned("bhc")
    prim = s.analyses["primary"]
    assert prim["status"] == "ok" and prim["execution"]["diagnostics_missing"] is True


def _standard_run(p, agents, model, extra_code=()):
    """Curate, plan (multivariable primary + Table 1), confirm, run the standard sequence."""
    from tests.test_sandbox import GOOD_WITH_DIAGNOSTICS
    agents.curate("bhc")
    model.script["plan"] = [{"analyses": [dict(PLAN_REPLY["analyses"][0], kind="multivariable")],
                             "table1": {"stratify_by": "gdm", "variables": ["age", "parity", "ethnicity"]}}]
    model.script["code"] = [fence(GOOD_WITH_DIAGNOSTICS)] + [fence(c) for c in extra_code]
    s = agents.propose_plan("bhc")
    for a in ("bhc", "coauthor"):
        s = p.confirm_plan(a, s.plan["sha"])
    return agents.run_planned("bhc")


def test_publication_tables_built_by_code_and_exported_to_word(setup):
    import io
    from docx import Document
    from saksi import tables
    p, agents, model, _ = setup
    s = _standard_run(p, agents, model)
    ts = tables.build(s, p.blobs)
    assert [t["analysis_id"] for t in ts] == ["s_table1", "primary"] and [t["number"] for t in ts] == [1, 2]
    t1, t2 = ts
    assert t1["columns"][0] == "Characteristic" and t1["columns"][-1] == "P value" and "(n = 20)" in t1["columns"][3]
    assert any(r.get("indent") for r in t1["rows"])                       # categories under their variable
    assert t2["columns"] == ["Variable", "Adjusted OR (95% CI)", "P value"]
    bmi = next(r for r in t2["rows"] if r["cells"][0] in ("booking_bmi", "BMI at booking"))
    assert re.fullmatch(r"\d\.\d\d \(\d\.\d\d to \d\.\d\d\)", bmi["cells"][1])
    notes = " ".join(t2["footnotes"])
    assert "events per predictor parameter" in notes and "confirmatory" not in notes and "xploratory" not in notes
    assert "Below the protocol's planned sample" in notes and "OR, odds ratio" in notes
    assert "Table 2 [A:primary]" in tables.listing(s, p.blobs)
    doc = Document(io.BytesIO(tables.to_docx(ts)))
    text = "\n".join(par.text for par in doc.paragraphs)
    assert "Table 1. Characteristics of the participants" in text and len(doc.tables) == 2
    assert doc.tables[1].rows[0].cells[1].text == "Adjusted OR (95% CI)"


def test_redo_replaces_the_run_keeps_the_label_and_the_record(setup):
    from tests.test_sandbox import GOOD_WITH_DIAGNOSTICS
    p, agents, model, _ = setup
    s = _standard_run(p, agents, model, extra_code=[GOOD_WITH_DIAGNOSTICS])
    old_figs = [f["number"] for f in s.figures if f["analysis_id"] == "primary"]
    with pytest.raises(WorkflowError, match="say what is wrong"):
        agents.redo("bhc", "primary", " ")
    s = agents.redo("bhc", "primary", "Parity should be categorical, as the protocol says.")
    new = s.analyses["primary.r1"]
    assert s.analyses["primary"]["superseded_by"] == "primary.r1" and new["supersedes"] == "primary"
    assert new["label"] == s.analyses["primary"]["label"] and new["plan_analysis_id"] == "primary"
    assert new["redo_reason"].startswith("Parity") and new["requested_by"] == "bhc"
    assert "THE AUTHORS ASKED FOR THIS ANALYSIS TO BE REDONE" in model.bodies[-1]["messages"][1]["content"]
    # The replaced run is out of the paper: not required, its figures unnumbered; the new one takes its place.
    from saksi.workflow import current
    assert "primary" not in current(s) and "primary.r1" in current(s)
    assert all(f["number"] is None for f in s.figures if f["analysis_id"] == "primary")
    assert sorted(f["number"] for f in s.figures if f["analysis_id"] == "primary.r1") == sorted(old_figs)
    r = p.check_text("There were 4.5 outcome events per predictor parameter [A:primary.r1]. "
                     "Table 1 describes the sample [A:s_table1]. Records were checked [A:s_data_quality].")
    assert r["analyses_not_reported"] == []
    with pytest.raises(WorkflowError, match="not a current analysis"):
        agents.redo("bhc", "primary", "again")


def test_variant_is_exploratory_and_the_original_stays(setup):
    from tests.test_sandbox import GOOD_WITH_DIAGNOSTICS
    p, agents, model, _ = setup
    _standard_run(p, agents, model, extra_code=[GOOD_WITH_DIAGNOSTICS])
    s = agents.run_request("bhc", "q1", "add waist circumference", based_on="primary")
    v = s.analyses["x_q1"]
    assert v["label"] == "exploratory" and v["basis"] == "post hoc" and v["kind"] == "multivariable"
    assert v["spec"]["based_on"]["analysis_id"] == "primary" and "BMI and GDM: add waist" in v["spec"]["title"]
    assert "THIS IS A VARIANT" in model.bodies[-1]["messages"][1]["content"]
    assert not s.analyses["primary"].get("superseded_by")                 # the original is still reported
    with pytest.raises(WorkflowError, match="not a current analysis"):
        agents.run_request("bhc", "q2", "x", based_on="s_table1")


def test_authors_change_table1_with_a_reason(setup):
    p, agents, model, _ = setup
    s = _standard_run(p, agents, model)
    with pytest.raises(WorkflowError, match="not columns in the dataset"):
        agents.change_table1("bhc", {"stratify_by": "gdm", "variables": ["age", "nope"]}, "fewer rows")
    with pytest.raises(WorkflowError, match="say why"):
        agents.change_table1("bhc", {"stratify_by": "gdm", "variables": ["age"]}, "")
    s = agents.change_table1("bhc", {"stratify_by": "gdm", "variables": ["age", "booking_bmi"]}, "Add BMI.")
    new = s.analyses["s_table1.r1"]
    assert new["supersedes"] == "s_table1" and new["by"] == "system" and new["spec"]["table1"]["variables"] == ["age", "booking_bmi"]
    t1 = p.blobs.get_json(new["outputs_sha256"])
    assert [r["variable"] for r in t1["rows"]] == ["age", "booking_bmi"]


def test_authors_reword_tables_on_the_draft_but_numbers_are_locked(setup):
    from saksi import tables
    p, agents, model, _ = setup
    bmi = _to_drafting(p, agents, model)
    model.script["write"] = [_draft(bmi, extra="", exploratory="An exploratory model gave a similar estimate [A:x_q1].")]
    s = agents.write_draft("bhc")
    text = p.blobs.get_text(s.drafts[1]["text_sha256"])
    # Every table and figure must be referred to in the text.
    bare = text.replace(REFS, "")
    r = p.check_text(bare)
    assert r["tables_not_referred"] == [1] and r["figures_not_referred"] and not r["ok"]   # x_q1 repeats Table 1
    assert p.check_text(text + "\n\nSee Table 9.")["references_unknown"] == ["Table 9"]
    # Words can change; numbers cannot.
    with pytest.raises(WorkflowError, match="must keep its numbers"):
        p.edit_draft("bhc", 1, text, "x", table_edits={"primary": {"footnotes": {"0": "n = 999 analysed."}}})
    with pytest.raises(WorkflowError, match="no table"):
        p.edit_draft("bhc", 1, text, "x", table_edits={"nope": {"title": "x"}})
    s = p.edit_draft("bhc", 1, text, "Relabelled Table 1",
                     table_edits={"primary": {"title": "Booking BMI and gestational diabetes", "rows": {"0": "Booking BMI, per kg/m2"}}})
    d = s.drafts[2]
    base = tables.build(s, p.blobs)
    shown = tables.apply_edits(base, p.blobs.get_json(d["table_edits_sha256"]))
    assert shown[0]["title"] == "Booking BMI and gestational diabetes" and shown[0]["rows"][0]["cells"][0] == "Booking BMI, per kg/m2"
    assert shown[0]["rows"][0]["cells"][1:] == base[0]["rows"][0]["cells"][1:] and shown[0]["edited"] == ["rows", "title"]
    # Carried forward by a text-only edit; an edit that changes nothing is refused.
    s = p.edit_draft("bhc", 2, text + "\n\nAdded a sentence.", "more")
    assert s.drafts[3]["table_edits_sha256"] == d["table_edits_sha256"]
    with pytest.raises(WorkflowError, match="nothing changed"):
        p.edit_draft("bhc", 3, p.blobs.get_text(s.drafts[3]["text_sha256"]), "same")
    # A heading keeps its group size; the group's code can become words.
    t = {"analysis_id": "a", "number": 1, "title": "T", "columns": ["Characteristic", "GDM: 0\n(n = 380)"],
         "rows": [{"cells": ["0", "1 (2.0)"]}], "footnotes": ["n = 5."]}
    assert tables.check_edits([t], {"a": {"columns": {"1": "No GDM\n(n = 380)"}, "rows": {"0": "Nulliparous"}}})
    with pytest.raises(ValueError, match="must keep its numbers"):
        tables.check_edits([t], {"a": {"columns": {"1": "No GDM"}}})


def test_each_analysis_gets_text_that_reports_it_and_is_number_checked(setup):
    from saksi import tables
    p, agents, model, _ = setup
    s = _standard_run(p, agents, model)
    s = agents.describe("bhc", "s_data_quality")
    d = s.analyses["s_data_quality"]["description"]
    assert d["section"] == "methods" and d["checks"]["ok"] and d["requested_by"] == "bhc" and d["by"] == "agent:writer"
    s = agents.describe("bhc", "primary")
    d = s.analyses["primary"]["description"]
    text = p.blobs.get_text(d["text_sha256"])
    assert d["section"] == "results" and d["checks"]["ok"] and "[[Table]]" in text and "[[Figure:" in text
    shown = tables.resolve(text, "primary", tables.build(s, p.blobs), s.figures)
    assert "Table 2" in shown and "Figure 1" in shown and "[[" not in shown
    # A number that is not in the results is sent back once; still wrong, it is kept with its findings.
    model.script["describe"] = ["The adjusted odds ratio was 9.99 [A:primary].",
                                "The adjusted odds ratio was 9.98 [A:primary]."]
    s = agents.describe("bhc", "primary", instruction="Shorter, please")
    d = s.analyses["primary"]["description"]
    assert [x["passed_check"] for x in d["attempts"]] == [False, False] and not d["checks"]["ok"]
    assert d["instruction"] == "Shorter, please" and "CHECK FINDINGS" in model.bodies[-1]["messages"][1]["content"]
    # The full draft gets the texts, tokens resolved.
    from saksi.agents import _texts
    from saksi.workflow import current
    assert "[METHODS] s_data_quality" in _texts(s, current(s), p.blobs)


def test_older_plans_get_a_place_in_the_sequence_by_rule_and_can_build_table1(setup):
    p, agents, model, _ = setup
    agents.curate("bhc")
    legacy = {k: v for k, v in PLAN_REPLY["analyses"][0].items() if k != "kind"}
    model.script["plan"] = [{"analyses": [legacy]}]
    s = agents.propose_plan("bhc")
    assert s.plan["analyses"]["primary"]["kind"] == "multivariable"       # inferred from the method
    for a in ("bhc", "coauthor"):
        s = p.confirm_plan(a, s.plan["sha"])
    assert not s.plan_locked.get("table1")
    s = agents.change_table1("bhc", {"stratify_by": "gdm", "variables": ["age", "parity"]})
    assert s.analyses["s_table1"]["status"] == "ok" and not s.analyses["s_table1"].get("supersedes")


def test_an_exploratory_analysis_can_be_left_out_with_a_reason_but_not_a_planned_one(setup):
    from saksi.workflow import current
    p, agents, model, _ = setup
    _standard_run(p, agents, model, extra_code=[GOOD])
    s = agents.run_request("bhc", "q1", "A characteristics table by GDM")
    with pytest.raises(WorkflowError, match="only an exploratory analysis added after the plan"):
        p.withdraw_analysis("bhc", "primary", "not wanted")
    with pytest.raises(WorkflowError, match="say why"):
        p.withdraw_analysis("bhc", "x_q1", "")
    before = s.analysis_set_sha()
    s = p.withdraw_analysis("bhc", "x_q1", "Superseded by Table 1")
    assert "x_q1" not in current(s) and s.analyses["x_q1"]["withdrawn"]["reason"] == "Superseded by Table 1"
    assert s.analysis_set_sha() != before                                  # acceptance must be given again
    assert all(f["number"] is None for f in s.figures if f["analysis_id"] == "x_q1")


def test_identical_tables_are_numbered_once_and_suggested_for_leaving_out(setup):
    from saksi import advisor, tables
    from tests.test_sandbox import GOOD_WITH_DIAGNOSTICS
    p, agents, model, _ = setup
    _standard_run(p, agents, model, extra_code=[GOOD_WITH_DIAGNOSTICS])
    s = agents.run_request("bhc", "q1", "the same model again", based_on="primary")
    ts = tables.build(s, p.blobs, with_duplicates=True)
    dup = next(t for t in ts if t["analysis_id"] == "x_q1")
    assert dup["number"] is None and dup["duplicate_of"] == {"analysis_id": "primary", "number": 2}
    assert [t["number"] for t in tables.build(s, p.blobs)] == [1, 2]
    sug = {x["id"]: x for x in advisor.suggestions(s, p.blobs)}
    assert sug["dup:x_q1"]["action"] == {"type": "leave_out", "analysis_id": "x_q1", "instruction": "Same results as primary"}


def test_proactive_suggestions_follow_epidemiological_rules(setup):
    from saksi import advisor
    p, agents, model, _ = setup
    s = _standard_run(p, agents, model)
    sug = {x["id"]: x for x in advisor.suggestions(s, p.blobs)}
    assert sug["epv:primary"]["action"]["type"] == "variant" and "Firth" in sug["epv:primary"]["action"]["instruction"]
    assert sug["crude:primary"]["action"]["type"] == "redo"               # the fixture reports adjusted only
    assert sug["text:primary"]["action"]["type"] == "rewrite_text"        # no text yet
    assert sug["c:primary"]["kind"] == "note"                             # C 0.65, not a prediction study
    assert "table1" not in sug
    # Once the sensitivity analysis exists, it is not suggested again.
    from tests.test_sandbox import GOOD_WITH_DIAGNOSTICS
    model.script["code"] = [fence(GOOD_WITH_DIAGNOSTICS)]
    s = agents.run_request("bhc", "q1", sug["epv:primary"]["action"]["instruction"], based_on="primary")
    assert "epv:primary" not in {x["id"] for x in advisor.suggestions(s, p.blobs)}


def test_the_assistant_answers_and_proposes_one_checked_action(setup):
    p, agents, model, _ = setup
    s = _standard_run(p, agents, model)
    s, prop = agents.assist("bhc", "Add waist circumference as a covariate", "primary")
    assert prop["action"] == {"type": "variant", "analysis_id": "primary", "instruction": "Add waist circumference as a covariate"}
    assert prop["reply"] and s.events_count if hasattr(s, "events_count") else True
    assert any(e.type == "llm.called" and e.payload.get("purpose") == "assist" for e in p.log.events())
    # It cannot leave out a planned analysis, whatever it proposes.
    model.script["assist"] = [{"reply": "Leaving it out.", "action": {"type": "leave_out", "analysis_id": "primary"}}]
    _, prop = agents.assist("bhc", "drop the primary analysis", "primary")
    assert prop["action"]["type"] == "none" and "planned analysis is reported" in prop["reply"]


def test_plan_states_the_report_aim_and_design_and_saksi_helps_settle_it(setup):
    p, agents, model, _ = setup
    agents.curate("bhc")
    model.script["plan"] = [dict(PLAN_REPLY, report={"aim": "Is booking BMI associated with GDM?", "design": "cohort",
                                                     "questions": ["Which objective is this paper's main question?"]})]
    s = agents.propose_plan("bhc")
    assert s.plan["report"] == {"aim": "Is booking BMI associated with GDM?", "design": "cohort",
                                "questions": ["Which objective is this paper's main question?"]}
    s, out = agents.plan_assist("bhc", "Only the primary objective", history=[{"who": "ai", "text": "Which objective?"}])
    assert out["redraft"] == "This paper reports the primary objective only." and out["questions"]
    assert "CONVERSATION SO FAR:\nSAKSI: Which objective?" in model.bodies[-1]["messages"][1]["content"]
    for a in ("bhc", "coauthor"):
        s = p.confirm_plan(a, s.plan["sha"])
    assert s.plan_locked["report"]["design"] == "cohort"


def test_draft_is_checked_against_the_guideline_for_the_design(setup):
    from saksi import guidelines
    p, agents, model, _ = setup
    bmi = _to_drafting(p, agents, model)
    model.script["write"] = [_draft(bmi, extra="", exploratory="An exploratory model gave a similar estimate [A:x_q1].")]
    s = agents.write_draft("bhc")
    assert "REPORTING GUIDELINE: STROBE" in model.bodies[-1]["messages"][1]["content"]
    s = agents.checklist("bhc")
    c = s.drafts[1]["checklist"]
    items = p.blobs.get_json(c["items_sha256"])
    assert c["guideline"].startswith("STROBE") and len(items) == len(guidelines.item_list(None))
    assert items[0]["status"] == "missing"
    assert items[1]["status"] == "reported" and items[1]["quote"]            # quotation found in the draft
    assert items[2]["status"] == "partly" and "not found" in items[2]["note"]  # an invented quotation is caught


def test_authors_edit_the_result_text_in_place_under_their_name(setup):
    p, agents, model, _ = setup
    s = _standard_run(p, agents, model)
    s = agents.describe("bhc", "primary")
    assert s.analyses["primary"]["description"]["heading"] == "BMI and GDM"          # the subheading line is parsed
    assert not p.blobs.get_text(s.analyses["primary"]["description"]["text_sha256"]).startswith("Heading")
    text = "In this exploratory analysis, [[Table]] gives the estimates; 4.5 events per parameter [A:primary]."
    s = p.edit_description("coauthor", "primary", text, heading="Booking BMI and gestational diabetes")
    d = s.analyses["primary"]["description"]
    assert d["by"] == "coauthor" and d["edited_by_author"] and d["checks"]["ok"] and d["heading"].startswith("Booking")
    with pytest.raises(WorkflowError, match="nothing changed"):
        p.edit_description("coauthor", "primary", text, heading="Booking BMI and gestational diabetes")
    s = p.edit_description("bhc", "primary", text.replace("4.5", "9.9"))
    assert s.analyses["primary"]["description"]["checks"]["untraced"] == 1             # still checked


def test_spss_missing_codes_are_missing_values_not_a_group():
    from saksi import standard
    from saksi.sandbox import LocalExecutor
    csv = "gdm,age,smoker\n0,30,1\n1,31,0\n#NULL!,29,#NULL!\n0,33,0\n1,35,1\n".encode()
    r = LocalExecutor().run(standard.code("table1", {"stratify_by": "gdm", "variables": ["age", "smoker"]}), csv, seed=1)
    assert [g["level"] for g in r.results["groups"]] == ["0", "1"]
    assert r.results["n_excluded"] == 1
